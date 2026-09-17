from pathlib import Path
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

root=Path('D:/AI Information Heirarchy/Heutrix')
work=root/'tools/org-structure-work'
out=root/'01-company/founders-and-roles/organisation-structure-2026-09/Heutrix-Organisational-Structure.docx'
navy='#033862'; teal='#01647C'
im=Image.new('RGB',(1800,1080),'white'); dr=ImageDraw.Draw(im)
def font(n,b=False): return ImageFont.truetype('C:/Windows/Fonts/'+('arialbd.ttf' if b else 'arial.ttf'),n)
def box(x,y,w,h,lines,dark=False):
    dr.rounded_rectangle((x,y,x+w,y+h),radius=20,fill=navy if dark else '#EFFAFA',outline='#CCDCDD',width=2)
    yy=y+24
    for text,size,b in lines:
        dr.text((x+w/2,yy),text,font=font(size,b),fill='white' if dark else navy,anchor='mt'); yy+=size+14
box(440,20,920,140,[('FOUNDER LEADERSHIP TEAM',30,True),('Janith  |  Rochelle  |  Sachin  |  Shashane',28,False)],True)
dr.line((900,160,900,220),fill=teal,width=5)
box(440,220,920,190,[('JANITH',36,True),('Managing Director',30,True),('Strategy, commercial and workflow solutions',27,False)])
dr.line((900,410,900,490),fill=teal,width=5)
dr.line((300,490,1500,490),fill=teal,width=5)
for x in [300,900,1500]: dr.line((x,490,x,550),fill=teal,width=5)
box(20,550,560,235,[('ROCHELLE',34,True),('Head of Operations',28,True),('and Client Delivery',28,True),('Projects  |  Adoption  |  Administration',23,False)])
box(620,550,560,235,[('SACHIN',34,True),('Head of Data, Insights',28,True),('and Business Assurance',28,True),('Analysis  |  Finance  |  Governance',23,False)])
box(1220,550,560,235,[('SHASHANE',34,True),('Head of Engineering',28,True),('and Technology',28,True),('Build  |  Infrastructure  |  Support',23,False)])
dr.text((900,860),'Client engagements coordinated by Rochelle',font=font(29,True),fill=navy,anchor='mt')
dr.text((900,910),'All four founders contribute within their specialist responsibilities',font=font(27),fill='#334155',anchor='mt')
dr.text((900,990),'Proposed operating relationships   •   Corporate authority remains separate',font=font(24),fill='#64748B',anchor='mt')
im.save(work/'org-chart.png')

d=Document(); sec=d.sections[0]
sec.page_height=Inches(11.7); sec.page_width=Inches(8.3)
sec.top_margin=Inches(.65); sec.bottom_margin=Inches(.6); sec.left_margin=sec.right_margin=Inches(.7)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2']:
    s=d.styles[name]; s.font.name='Arial'; s.font.color.rgb=RGBColor(0,0,0)
d.styles['Normal'].font.size=Pt(10)
d.styles['Normal'].paragraph_format.space_after=Pt(7)
d.styles['Normal'].paragraph_format.line_spacing=1.08
d.styles['Title'].font.size=Pt(29)
d.styles['Heading 1'].font.size=Pt(20)
d.styles['Heading 2'].font.size=Pt(12)
foot=sec.footer.paragraphs[0]; foot.text='HEUTRIX  |  Proposed organisational structure  |  16 September 2026'; foot.style='Caption'
foot.runs[0].font.size=Pt(8)
foot.runs[0].font.color.rgb=RGBColor.from_string('64748B')
for element in d.styles.element.iter():
    for child in list(element):
        if child.tag == qn('w:pBdr'): element.remove(child)
def p(t,style=None): return d.add_paragraph(t,style)
def h(t): d.add_heading(t,2)
def page(title): d.add_page_break(); d.add_heading(title,1)
def table(headers,rows,widths):
    t=d.add_table(rows=1,cols=len(headers)); t.autofit=False
    for c,w in zip(t.columns,widths): c.width=Inches(w)
    for c,txt in zip(t.rows[0].cells,headers): c.text=txt
    for row in rows:
        for c,txt in zip(t.add_row().cells,row): c.text=txt
    for ri,row in enumerate(t.rows):
        for ci,c in enumerate(row.cells):
            c.width=Inches(widths[ci]); pr=c._tc.get_or_add_tcPr()
            sh=OxmlElement('w:shd'); sh.set(qn('w:fill'),'033862' if ri==0 else ('F2F7F8' if ri%2 else 'FFFFFF')); pr.append(sh)
            borders=OxmlElement('w:tcBorders')
            for edge in ['top','left','bottom','right']:
                e=OxmlElement('w:'+edge); e.set(qn('w:val'),'single'); e.set(qn('w:sz'),'4'); e.set(qn('w:color'),'D9D9D9'); borders.append(e)
            pr.append(borders)
            mar=OxmlElement('w:tcMar')
            for edge in ['top','left','bottom','right']:
                e=OxmlElement('w:'+edge); e.set(qn('w:w'),'95'); e.set(qn('w:type'),'dxa'); mar.append(e)
            pr.append(mar)
            for pp in c.paragraphs:
                pp.paragraph_format.space_after=Pt(2); pp.paragraph_format.space_before=Pt(2)
                for r in pp.runs:
                    r.font.size=Pt(9); r.bold=ri==0
                    if ri==0:r.font.color.rgb=RGBColor(255,255,255)
        trpr=row._tr.get_or_add_trPr(); trpr.append(OxmlElement('w:cantSplit'))
    t.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
    p('')

d.add_picture(str(root/'06-marketing/brand/logo-2026-09/heutrix-logo-colour.png'),width=Inches(1.7))
d.add_heading('Heutrix organisational structure',0)
p('Founder responsibilities and operating areas','Subtitle')
p('Proposed for founder review  |  16 September 2026')
p('Heutrix should operate through four complementary founder functions, with Janith coordinating the business and Rochelle coordinating client delivery. Each business area has one accountable lead, supported by the other founders where their expertise is needed.')
d.add_picture(str(work/'org-chart.png'),width=Inches(6.85))
h('Contents')
p('01  Organisation overview and reporting relationships\n02  Founder roles and responsibilities\n03  Business areas and folder contents\n04  Delivery responsibilities and decision making')
p('The structure is a recommendation. It does not appoint officers, change signing authority or assign permanent folder owners. Existing confirmed responsibilities remain in effect until an agreed update is recorded.')

page('Founder roles and responsibilities')
roles=[
('Janith','Managing Director — Strategy, Commercial and Workflow Solutions','Workflow implementation and experience helping build a disability support business.','Set business priorities and offer direction; lead fit calls, sales and partnerships; translate client needs into workflow designs; coordinate decisions across functions.','Clear scope, a qualified pipeline, viable commercial commitments and aligned founder priorities.'),
('Rochelle','Head of Operations and Client Delivery','Project delivery and disability-provider administration and operations.','Plan engagements and capacity; coordinate clients and milestones; maintain delivery methods; lead staff training and adoption; organise administration, invoicing and handover.','Predictable delivery, clear client communication, complete handovers and timely administrative follow-through.'),
('Sachin','Head of Data, Insights and Business Assurance','Data and analytics.','Lead diagnostic analysis and outcome measurement; maintain financial models and management reporting; coordinate data governance and risk records; support campaign measurement.','Reliable evidence, visible project economics and documented handling and review arrangements.'),
('Shashane','Head of Engineering and Technology','Backend engineering and infrastructure.','Design and build integrations and automation; maintain infrastructure and the website; manage technical access controls, testing, release readiness and support; retain confirmed enquiry cover.','Maintainable solutions, controlled releases, reliable infrastructure and responsive technical support.')]
for name,title,fit,scope,result in roles:
    h(name+'   '+title)
    p('Documented strengths: '+fit)
    p(scope)
    pp=p('Accountable outcome: '+result); pp.runs[0].italic=True
p('Finance and governance ownership means coordinating the work and obtaining appropriate specialist input. The documented founder summaries do not establish accounting, legal or privacy qualifications.')

page('Business areas and folder contents')
p('The folders remain the operating library. Proposed leads maintain their area and involve affected founders when changes alter scope, cost, delivery, public claims or information handling.')
table(['Area','Proposed lead','Contents and responsibilities'],[
('00 Control','Janith','Priorities, decisions, launch status, dependencies, risks and shared coordination.'),
('01 Company','Janith','Company records, founder roles, approved systems and insurance records. Rochelle supports administration.'),
('02 Market and offers','Janith','Target markets, positioning, service definitions, scope boundaries and private pricing policy.'),
('03 Sales','Janith','Enquiries, qualification, fit calls, proposals, pipeline, follow-up and delivery handover.'),
('04 Delivery','Rochelle','Delivery methods, project planning, workflow design, testing, training, acceptance and closeout.'),
('05 Legal privacy and risk','Sachin','Review coordination, data-handling standards, risk and incident records. Directors and advisers retain their respective authority.'),
('06 Marketing','Janith','Brand, website content, campaigns, resources and case studies. Sachin supports publishing and measurement.'),
('07 Finance','Sachin','Budgets, forecasts, costing, margins and reporting. Rochelle coordinates invoicing and payment follow-up.'),
('08 Client project template','Rochelle','Reusable engagement structure, project records, deliverables and handover templates. Live client records remain outside Git.'),
('Applications and tools','Shashane','Website, automation, integrations, infrastructure, technical documentation and maintenance.')
],[1.55,1,4.3])
h('Supporting relationships')
p('Rochelle checks delivery capacity; Sachin checks economics and evidence; Shashane checks technical feasibility. Janith brings these inputs together before a client commitment is made. Marketing claims and sales promises must stay aligned with the agreed offer and delivery capability.')

page('Delivery responsibilities and decision making')
h('How the offers draw on the team')
table(['Offer','Specialist contributions'],[
('Heutrix Diagnostics','Janith frames the workflow problem; Sachin leads analysis; Rochelle coordinates the engagement; Shashane assesses technical constraints.'),
('Heutrix Workflow Transformation','Rochelle leads delivery; Janith designs the workflow; Shashane implements technology; Sachin measures results.'),
('Heutrix AI Guardrails','Sachin coordinates governance requirements; Shashane implements technical controls; Rochelle leads procedures and adoption; Janith aligns scope with client needs.')
],[2,4.85])
h('One accountable lead for each decision')
table(['Decision','Proposed lead','Required input'],[
('Qualification and proposed scope or price','Janith','Rochelle on capacity; Sachin on economics; Shashane on feasibility.'),
('Delivery plan and internal closeout','Rochelle','Specialist readiness checks and the client’s agreed acceptance process.'),
('Financial reporting and margin visibility','Sachin','Rochelle on actuals and collections; Janith on commercial priorities.'),
('Technical release readiness','Shashane','Rochelle on delivery readiness; Sachin on data requirements.'),
('Privacy and incident coordination','Sachin','Shashane on technical response; Rochelle on operations; relevant advisers as needed.')
],[2.3,1,3.55])
h('Operating rhythm')
p('Weekly founder meeting: review leads, delivery capacity, cash, risks and decisions. Rochelle coordinates project check-ins with the specialists involved. Review financial performance and offer effectiveness monthly. Agree spending limits, escalation thresholds and backups when adopting the structure.')
h('Confirmed responsibilities and adoption')
p('Current records confirm Shashane on enquiries with Rochelle as backup, Janith on fit calls with Rochelle as backup, and Janith on merge coordination. Janith and Sachin are recorded as authorised director signatories; execution arrangements remain document-specific. Other allocations here are proposed.')
p('To adopt: agree titles, capacity, decision boundaries and backups, then date the agreed assignments in the role register. Use Shashane consistently here; confirm the existing Sashane spelling discrepancy before updating the register.')
p('Source basis: founder role register; workspace README; verified facts and dependency registers; current Heutrix brand guide. These records support the founder strengths and business areas; the organisation design is a recommendation.')
d.save(out)
print(out)
