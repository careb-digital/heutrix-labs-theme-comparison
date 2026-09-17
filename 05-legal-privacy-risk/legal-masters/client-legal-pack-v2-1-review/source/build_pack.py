from pathlib import Path
from copy import deepcopy
from zipfile import ZipFile, ZIP_DEFLATED
from lxml import etree
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
import re, json, hashlib

PACK = Path(__file__).resolve().parents[1]
ROOT = PACK.parents[1]
OLD = PACK.parent / 'client-legal-pack-v1' / 'Word'
OUT = PACK / 'Word'
OUT.mkdir(parents=True, exist_ok=True)
REF = Path('C:/Users/CareBest/.codex/plugins/cache/openai-curated-remote/openai-templates/0.1.1/skills/artifact-template-legal-memorandum/assets/reference.docx')
IDENTITY = 'Heutrix Pty Ltd | ABN 64 702 109 662 | ACN 702 109 662'
VERSION = '2.1 review draft | 13 September 2026'
STATUS = 'REVIEW DRAFT - DO NOT ISSUE OR SIGN'
changes = []

def all_paras(d):
    yield from d.paragraphs
    def cell_paras(cell):
        yield from cell.paragraphs
        for t in cell.tables:
            for row in t.rows:
                for c in row.cells: yield from cell_paras(c)
    seen=set()
    for t in d.tables:
        for row in t.rows:
            for c in row.cells:
                if c._tc not in seen:
                    seen.add(c._tc)
                    yield from cell_paras(c)
    for s in d.sections:
        for part in (s.header,s.footer,s.first_page_header,s.first_page_footer):
            yield from part.paragraphs
            for t in part.tables:
                for row in t.rows:
                    for c in row.cells: yield from cell_paras(c)

def set_text(p,text):
    if p.runs:
        p.runs[0].text=text
        for r in p.runs[1:]: r.text=''
    else: p.add_run(text)

def paragraph_before(p,text):
    n=p.insert_paragraph_before(text)
    n.style='Normal'
    return n

def apply_style(d,short,internal=False):
    d.settings.odd_and_even_pages_header_footer=False
    for s in d.styles:
        if s.type==1 and (s.name.startswith('Heading') or s.name in ('Title','Subtitle','Header')):
            s.font.color.rgb=RGBColor(0,0,0)
        if s.name=='Title':
            for border in list(s.element.xpath('./w:pPr/w:pBdr')):border.getparent().remove(border)
    for p in all_paras(d):
        if p.text.startswith('Each signatory confirms'):
            p.paragraph_format.keep_with_next=True
        if p.style.name.startswith('Heading') or p.style.name in ('Title','Subtitle','Header'):
            for r in p.runs:
                r.font.color.rgb=RGBColor(0,0,0)
                r.font.underline=False
        if p.style.name=='Title':
            borders=p._p.pPr.find(qn('w:pBdr')) if p._p.pPr is not None else None
            if borders is not None: borders.getparent().remove(borders)
    for s in d.sections:
        s.different_first_page_header_footer=False
        s.header.paragraphs[0].text = f'HEUTRIX | {short} | ' + ('INTERNAL REVIEW' if internal else STATUS)
        for r in s.header.paragraphs[0].runs:
            r.font.size=Pt(8); r.font.color.rgb=RGBColor(0,0,0)
        s.footer.paragraphs[0].text = 'ABN 64 702 109 662 | ACN 702 109 662 | v2.1 review | Page '
        fld=OxmlElement('w:fldSimple'); fld.set(qn('w:instr'),'PAGE'); s.footer.paragraphs[0]._p.append(fld)
        for r in s.footer.paragraphs[0].runs: r.font.size=Pt(8)
    for t in d.tables:
        pr=t._tbl.tblPr
        borders=pr.find(qn('w:tblBorders'))
        if borders is None: borders=OxmlElement('w:tblBorders'); pr.append(borders)
        for side in ['top','left','bottom','right','insideH','insideV']:
            x=OxmlElement('w:'+side)
            for k,v in [('val','single'),('sz','4'),('color','D9D9D9')]: x.set(qn('w:'+k),v)
            borders.append(x)
        for row in t.rows:
            for trh in list(row._tr.xpath('./w:trPr/w:trHeight')):
                trh.set(qn('w:hRule'),'atLeast')
    d.core_properties.author='Heutrix'
    d.core_properties.last_modified_by='Heutrix'
    d.core_properties.comments='Controlled review draft. Approval is recorded separately.'
    d.core_properties.version='2.1 review'

generic = [
 ('governed by Victoria, Australia', 'governed by the laws of Victoria, Australia'),
 ('Payable only if Heutrix is registered or required to be registered and issues a valid tax invoice','[COMPLETE verified GST status, any lawful GST and total payable; current source records non-registration]'),
 ('[  ] AI limited to public/general research with no Client Data or Confidential Information.','[  ] Public/general research only, with no Client Data or Confidential Information; attach completed form B identifying the approved tool, settings and reviewer.'),
 ('Heutrix Client Services Agreement v1.0, dated 1 August 2026','Heutrix Client Services Agreement v2.1 review draft, dated 13 September 2026; replace with exact approved issue version before signature'),
 ('1.0 | 1 August 2026',VERSION),
 ('ACN / ABN: ______________________________','ABN 64 702 109 662 | ACN 702 109 662'),
 ('________________________________  COMPLETE BEFORE ISSUE','ABN 64 702 109 662 | ACN 702 109 662'),
 ('Workflow Diagnostic','Heutrix Diagnostics'),
 ('workflow builds, visibility builds, Safe AI packs and tailored internal systems','Heutrix Workflow Transformation and Heutrix AI Guardrails'),
 ('Client requirement / accepted gap','Client requirement and Heutrix approval reference'),
 ('Five business days unless stated here: ____________________','[APPROVE review period in business days]'),
 ('Amount ex GST','Amount before any lawful GST'),
 ('Total fees ex GST','Total fee before any lawful GST'),
 ('Fees ex GST','Fee change before any lawful GST'),
]

prefixes = {
 'A SOW may add or change terms': 'A SOW may add project-specific terms. A variation to this agreement must identify the clause changed and be signed by authorised representatives. Subject to non-excludable law, the order for an expressly agreed inconsistency is: signed Special Conditions identifying the variation; signed SOW; activated data or AI schedule; this agreement. A general scope description does not override an express data prohibition or approve AI. A purchase order is administrative only unless both parties sign an amendment.',
 'Heutrix provides operational workflow': 'Heutrix provides Heutrix Diagnostics, Heutrix Workflow Transformation and Heutrix AI Guardrails for non-clinical operational work. Services exclude legal, tax, financial, employment, clinical, medical, regulatory, official audit, NDIS registration readiness, mock audit, compliance assurance, certification and medical-device advice. This pack does not establish capability or authority to provide regulated services.',
 'The Client will review each Deliverable': 'The Client will review each Deliverable against the Acceptance Criteria within [APPROVE review period] business days as stated in the SOW, and give written acceptance or specific evidence of non-conformity. Heutrix will correct an in-scope material non-conformity within a reasonable agreed period and resubmit it. The SOW states included refinement rounds and the separate defect procedure.',
 'If the Client does not respond within': 'If the Client does not respond by the agreed review date, Heutrix will remind the Client and escalate to the nominated decision-maker. Silence does not constitute acceptance. The parties will agree the resulting schedule and any properly payable milestone under the signed SOW. Acceptance does not waive latent defects, express warranties, consumer guarantees or non-excludable rights.',
 'The Client will pay the fees': 'The Client will pay only the fees and approved expenses in the signed SOW. Payment milestones, invoice triggers, advance payments and the due period of [APPROVE payment period] calendar days must be completed before signature. No public price, automatic Diagnostics fee credit or unapproved expense is incorporated.',
 'Amounts are exclusive of GST': 'The SOW must state the base fee, any lawfully payable GST and total payable. The verified 8 September 2026 record states Heutrix is not currently registered for GST. Recheck status before issue: do not add GST or label an invoice a tax invoice while not registered. If registration or a taxable supply treatment changes, apply the reviewed contractual adjustment and valid invoice requirements; no automatic tax uplift is approved by this draft.',
 'The Client may dispute an invoice': 'The Client may dispute an invoice in good faith by giving reasons promptly, using the process and response period [APPROVE] stated in the SOW. It must pay the undisputed amount by the due date. The parties will work promptly to resolve the disputed part. Any recovery costs or interest require an expressly agreed, legally reviewed and proportionate term; none is added by default.',
 'The SOW must select the Deliverable IP model': 'The SOW must select and fully describe the Deliverable IP model before signature. No model is selected by default. Where the SOW selects client ownership, Heutrix assigns its rights in project-specific Deliverables on the agreed payment condition, excluding Background Materials and third-party materials. The scope identifies source/editable exports and any further assignment steps and costs.',
 'These obligations continue for five years': 'These obligations continue for [APPROVE confidentiality survival period] after disclosure, and for trade secrets and Personal Information for as long as the information remains protected by law or confidential in nature.',
 'The parties will minimise Personal Information': 'The parties will minimise Personal Information and use synthetic or appropriately de-identified examples where practicable. Client records, sensitive information and credentials must not be placed in public enquiry forms, ordinary project messages, public AI services or unapproved repositories. Basic business-contact details may use the approved enquiry route under the published website policy.',
 "For all claims arising from an affected SOW": "For all claims arising from an affected SOW, each party's aggregate general liability is capped at [APPROVE general cap basis and amount]. Complete the reviewed cap in the final agreement before signature; no amount is selected by this draft.",
 "For claims arising from breach of confidentiality": "For claims arising from breach of confidentiality, privacy or security obligations, or the IP indemnity, each party's aggregate liability is capped at [APPROVE higher cap basis and amount]. Review scope, exclusions and insurance alignment before signature.",
 'Each party will maintain insurance required': 'Each party will maintain insurance required by law and any insurance expressly agreed in the SOW. Before commencement, record Heutrix\'s verified applicable cover and approval evidence. A client acknowledgment of missing cover does not replace Heutrix\'s insurance review and onboarding gate. Provide evidence of stated cover through a restricted channel on reasonable request.',
 'Heutrix may suspend only the affected Services': 'Heutrix may suspend only the affected Services if an undisputed amount remains overdue after the written notice and cure period of [APPROVE] business days, or continuing creates a material security, privacy, safety or legal risk. Give reasons, limit suspension, preserve Client Data and restore Services promptly when the issue is resolved. No charge arises for avoidable delay caused by an unjustified suspension. Specific pause, reschedule and restart effects must be agreed in the SOW.',
 'Either party may terminate this agreement': 'Either party may terminate this agreement or an affected SOW for material breach not cured within [APPROVE breach cure period] business days after written notice, immediately for an incurable material breach, or where the other party becomes insolvent to the extent permitted by law.',
 'Either party may terminate a SOW for convenience': 'Either party may terminate a SOW for convenience on [APPROVE notice period] business days\' notice. The Client pays for conforming Services performed and approved unavoidable commitments up to termination. Heutrix refunds prepaid amounts for Services not performed and provides reasonable handover of completed paid work. Review the cancellation/refund calculation and any agreed exceptions before issue.',
 'A notice must be sent to the notice email': 'A notice must be sent to the verified notice email and address stated in the signing particulars, or a replacement notified in writing. Record an acknowledgment or other reliable evidence of receipt; a bounced message is not receipt. Apply [APPROVE receipt and business-hours rule] in the final agreement. Ordinary project communications do not change the notice particulars or confer signing authority.',
 'Use only the approved Client tenant': 'Use only the exact Client tenant or Heutrix workspace approved for the data class in the SOW and form A. No platform category, including Google Workspace or Microsoft 365/SharePoint, is approved merely by this draft. Record applications, administrators, settings, locations and vendor approval.',
 'Heutrix will notify the Client without undue delay': 'Heutrix will notify the Client without undue delay after becoming aware of a Security Incident affecting Client Data, and within the expressly agreed notification period [APPROVE maximum and trigger] where applicable. Meet any faster legal obligation. Initial notice may be incomplete and must be updated as facts become available.',
 'At the earlier of instruction, access expiry': 'At a lawful instruction, access expiry or project close, Heutrix will return agreed Client Data and apply the approved category-specific retention/deletion schedule. The SOW and form A must state working-copy deletion, lawful retention/holds, backup expiry, responsible person and verification method. Retained copies remain protected. Form E records completion and any exception; no fixed deletion period is implied.',
 'The parties will avoid exchanging identifiable': 'The parties will avoid exchanging identifiable patient, participant, worker or other sensitive information during preliminary discussions. This NDA does not authorise that exchange. If Personal Information is necessary, pause and obtain a separately approved lawful handling arrangement before collection. Applicable Privacy Law means privacy and health-records laws binding on the relevant party or handling.',
 'On request or when discussions end': 'On request or when discussions end, stop using and return or securely delete Confidential Information within [APPROVE NDA return/deletion period] business days, subject to lawful retention or a documented protected backup exception. Record the exception, responsible person and review/expiry date. Retained copies remain protected and cannot be used for another purpose.',
 'This agreement applies to disclosures made': 'This agreement applies to disclosures made for [APPROVE NDA disclosure term] after the Effective date. Confidentiality and use restrictions continue for [APPROVE NDA survival period] after each disclosure; trade secrets and Personal Information remain protected for as long as confidential or protected by law.',
 'One remote structured discovery workshop': 'Discovery interviews/workshops: [APPROVE number, duration, format and included participants in the written scope].',
 'Review of up to five relevant documents': 'Document/evidence review: [APPROVE document or page limit and evidence depth], using synthetic or appropriately de-identified material unless a separate lawful data authorisation is approved.',
 'One remote findings presentation': 'Findings readout: [APPROVE number, duration, format and audience].',
 'The Client has five business days': 'The Client has [APPROVE review period] business days to accept the deliverable in writing or provide evidence against the criteria. Included correction/refinement rounds: [APPROVE]. New scope or preference changes require a signed Change Request. Silence is not acceptance under this draft.',
 'The Client provides agreed documents': 'The Client provides agreed safe documents and attendee availability by [COMPLETE agreed input date].',
 'Fixed fee: AUD 950': 'Private fixed fee: AUD [APPROVE fee]. GST treatment and total payable: [COMPLETE verified treatment and total]. Approved invoice milestones, due dates and payment/start condition: [COMPLETE]. No payment is requested under this review draft.',
 'The fixed fee covers only the Included Services': 'The fixed fee covers the completed Included Services and approved correction rounds only. Extra work and third-party or travel costs require prior written approval through a signed Change Request. Diagnostics is standalone paid work; no later purchase or automatic fee credit is required.',
 'The Client-owned Deliverable model in the CSA applies': 'Select the reviewed IP model before signature: [APPROVE client ownership / defined licence / other reviewed arrangement]. Identify the maps/report, editable formats, embedded Background Materials and third-party rights. The selected CSA/SOW terms apply; no IP model is selected by this blank order.',
 '[  ] Client-owned Deliverables. The CSA default': '[  ] Client-owned Deliverables. Complete the assignment/payment trigger and source/editable formats; exclude identified Background Materials and third-party materials. No default selection applies.',
}

def reconcile_native(num,newname):
    src=next(p for p in OLD.glob(f'{num:02}_*.docx') if not p.name.startswith('._'))
    d=Document(src)
    for p in all_paras(d):
        old=p.text; text=old
        for start,repl in prefixes.items():
            if text.startswith(start): text=repl; break
        if num==2:
            exact={
                'The workflow is sufficiently narrow to cover in one workshop and starts and ends at the points recorded below.':'The workflow is bounded by the start/end points below and can be investigated within the approved discovery effort.',
                'Before commencement, following a valid invoice':'[APPROVE milestone, due date and start condition]',
                'Not approved unless required for an approved redacted document':'Not approved unless specifically authorised in form A with a secure storage and deletion plan',
                'One workflow. A clear current-state view. A practical next step.':'Standalone paid workflow investigation and decision support',
                'AUD 950 plus GST, if GST is lawfully payable':'AUD [APPROVE fee] | [COMPLETE lawful GST treatment and total]',
                'Within 10 business days of inputs':'[COMPLETE agreed report target]',
                'Within 5 business days of draft':'[COMPLETE agreed readout target]',
                '10 business days after payment and receipt of agreed inputs':'[COMPLETE target and approved start prerequisites]',
                'Client-controlled Google Workspace / SharePoint / other: ____________________':'[APPROVE exact system, tenant and evidence reference]',
                'No Client Confidential Information or Client Data in AI. Public/general research only unless a signed AI approval is attached.':'No AI use is approved unless an applicable signed form B is attached.',
                'Return/delete project Client Data within 30 days after acceptance, subject to lawful retention and backup cycle':'[APPROVE category-specific return, retention, deletion, backup expiry and verification]',
            }
            text=exact.get(text,text)
        for oldpart,newpart in generic: text=text.replace(oldpart,newpart)
        if 'GST is registered' in text: pass
        if text!=old:
            set_text(p,text)
            changes.append({'document':newname,'before':old,'after':text})
    title = {'1':'Client Services Agreement','2':'Heutrix Diagnostics Order Form','3':'Statement of Work','4':'Mutual Confidentiality Agreement','5':'Project Control Forms'}[str(num)]
    set_text(d.paragraphs[1],title)
    d.paragraphs[1].style='Title'
    for p in d.paragraphs[:5]:
        for r in p.runs:r.font.color.rgb=RGBColor(0,0,0)
        if p._p.pPr is not None:
            for border in list(p._p.pPr.findall(qn('w:pBdr'))):p._p.pPr.remove(border)
    for t in list(d.tables):
        if len(t.rows)==1 and len(t.columns)==1:
            for p in t.cell(0,0).paragraphs:t._tbl.addprevious(deepcopy(p._p))
            t._tbl.getparent().remove(t._tbl)
    if num==1:
        for p in d.paragraphs:
            for br in list(p._p.xpath('.//w:br[@w:type="page"]')):br.getparent().remove(br)
            if p.text.startswith('Schedule 1') or p.text.startswith('Schedule 2'):
                p.paragraph_format.page_break_before=True
        for p in list(d.paragraphs):
            if not p.text.strip() and not p._p.xpath('.//w:drawing|.//w:sectPr'):
                p._p.getparent().remove(p._p)
    # Put the issue gate immediately before substantive numbered sections.
    anchor=next((p for p in d.paragraphs if p.text.startswith('1.')),d.paragraphs[3])
    paragraph_before(anchor,'REVIEW DRAFT. Complete document 09 approvals and every required field before client issue. No legal review, business approval or effective date is recorded by this file. Keep completed client copies outside this repository.')
    if num in (1,4):
        paragraph_before(anchor,'Heutrix business and notices address: 4 Ellaroo Ct, Clyde North VIC 3978, Australia. Owner-confirmed authorised signatories: Janith Dharmasinghe and Sachin Wickramasuriya, Directors. Formal notice email [APPROVE]; execution method and signatory selection for this document [COMPLETE AFTER REVIEW]. Client notice address and authorised signatory [COMPLETE].')
    if num==3:
        paragraph_before(anchor,'Product [SELECT Heutrix Workflow Transformation / Heutrix AI Guardrails]. If both are included, price and define each scope and acceptance separately. Diagnostics normally uses document 02. No purchase of Diagnostics or AI Guardrails is compulsory for Transformation.')
        paragraph_before(anchor,'Transformation baseline: one bounded non-clinical workflow; Diagnose, Triage, Build, Handover; conditional typical 2-4 weeks after agreed prerequisites. Include the 60-90-minute workshop, two review checkpoints, testing, guidance, training and handover within approved scope/capacity. State exact systems, roles, integrations and refinement rounds.')
        paragraph_before(anchor,'Transformation defect commitment: 30-day stabilisation and correction of qualifying in-scope defects. Agreed trigger [APPROVE]; start/end dates [COMPLETE]; reporting channel, remedy and exclusions [COMPLETE]. The period does not limit statutory rights. Enhancements and ongoing support are separately scoped.')
        paragraph_before(anchor,'AI Guardrails scope: specify teams, tools and use cases; approved governance outputs, assessment depth, scenarios, sponsor review, briefing and handover owner. These working requirements must be validated and scoped. No technical enforcement, continuous monitoring, certification or compliance outcome is implied. Any actual AI use requires form B.')
    if num==5:
        index_entry=next(p for p in d.paragraphs if p.text.startswith('E.') and 'Handover' in p.text)
        for entry in ['F. Privacy request and complaint record', 'G. Incident assessment and response record']:
            added=deepcopy(index_entry._p)
            index_entry._p.addnext(added)
            from docx.text.paragraph import Paragraph
            index_entry=Paragraph(added,index_entry._parent)
            set_text(index_entry,entry)
        for label,fields,body in [
            ('F Privacy request and complaint record',['Request reference and received date','Request type and safe contact reference','Identity or authority check and method','Information locations and client involvement','Applicable duty deadline and action owner','Assessment and lawful retention or refusal reasons','Response date channel and restricted evidence','Completion actions and follow up date'],'Internal restricted record. Do not place identity documents or sensitive request content in this repository. Use with SOP 10.'),
            ('G Incident assessment and response record',['Incident reference and discovery time with timezone','Reporter and incident lead with backup','Affected systems and information categories','Restricted evidence location and preservation steps','Containment actions owner and time','Applicable contractual and statutory deadlines','Serious harm and remediation assessment','Notification decision reasons reviewer and time','Notices sent evidence and next update time','Recovery approval residual actions and review date'],'Internal restricted record. Follow SOP 09 immediately. Do not wait for a complete form before containment or required notification. Do not enter raw personal data or secrets.'),
        ]:
            d.add_page_break(); d.add_heading(label,1); d.add_paragraph(body)
            for f in fields: d.add_paragraph(f+'\n[COMPLETE]',style='Normal')
        for p in d.paragraphs:
            if p.text.startswith('When signed, this Change Request'):
                p.insert_paragraph_before('Change effective date [COMPLETE] | replaced SOW clause/version [COMPLETE] | total fee and lawful GST effect [COMPLETE]')
            if p.text.startswith('Approval condition:'):
                p.insert_paragraph_before('A classification option cannot override a data prohibition. Attach form A where required; approve provider/retention/region changes before use.')
    apply_style(d,f'{num:02}')
    d.save(OUT/newname)

def clean_markdown(s):
    return s.replace('**','').replace('–','-').replace('—','-').replace('’',"'")

def new_document(source,filename,short,internal=False):
    d=Document()
    sec=d.sections[0]; sec.page_width=Inches(8.27); sec.page_height=Inches(11.69)
    sec.top_margin=Inches(.72);sec.bottom_margin=Inches(.7);sec.left_margin=Inches(.8);sec.right_margin=Inches(.8)
    normal=d.styles['Normal']; normal.font.name='Calibri'; normal.font.size=Pt(11)
    normal.paragraph_format.space_after=Pt(6);normal.paragraph_format.line_spacing=1.08
    if short=='08':
        normal.paragraph_format.space_after=Pt(5)
        normal.paragraph_format.line_spacing=1.04
    d.styles['Title'].font.size=Pt(23)
    for hn,sz in [('Heading 1',15),('Heading 2',12)]:
        d.styles[hn].font.size=Pt(sz);d.styles[hn].paragraph_format.space_before=Pt(12); d.styles[hn].paragraph_format.space_after=Pt(6)
    for line in source.read_text(encoding='utf-8').splitlines():
        line=clean_markdown(line.strip())
        if not line: continue
        if line.startswith('# '): d.add_paragraph(line[2:],style='Title')
        elif line.startswith('## '): d.add_heading(line[3:],1)
        elif line.startswith('### '): d.add_heading(line[4:],2)
        elif line.startswith('- '): d.add_paragraph(line[2:],style='List Bullet')
        else: d.add_paragraph(line)
    apply_style(d,short,internal)
    d.save(OUT/filename)

def memo():
    # Patch only text-bearing XML parts; retain template page geometry and all opaque parts.
    ns={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
    with ZipFile(REF) as zin: parts={n:zin.read(n) for n in zin.namelist()}
    tree=etree.fromstring(parts['word/document.xml']); ps=tree.findall('w:body/w:p',ns)
    text_by_index={
      1:'Heutrix launch readiness',3:'To:\tHeutrix founders and legal reviewer',4:'',5:'From:\tHeutrix launch preparation',6:'Date:\t13 September 2026',7:'Re:\tLegal pack and operating guidance for 14 September',
      10:'Background',
      11:'\tThe public website privacy policy and terms are approved and published. The client legal pack has been reconciled as version 2.1 for review. It cannot yet be issued for signature: formal notice email/execution, commercial, insurance, data and legal approvals remain open. This memo records preparation status, not legal advice or approval.',
      12:'\tHeutrix Pty Ltd is the documented entity, ABN 64 702 109 662 and ACN 702 109 662. The verified 8 September record states GST non-registration and VIC 3978. The owner has supplied 4 Ellaroo Ct, Clyde North VIC 3978, Australia as business and notices address, and confirmed Janith Dharmasinghe and Sachin Wickramasuriya, Directors, as authorised signatories. Formal notice email and execution method remain to be settled. Recheck changeable particulars before issue.',
      13:'The controlled sources for this package are:',
      14:'The legal review register and website PR 16 approval record;',
      15:'The registration, pricing, approved-systems and delivery records; and',
      16:'The current three-product offer and owner-validation records.',
      17:'\tUse Heutrix Diagnostics for standalone paid investigation, Heutrix Workflow Transformation for a bounded implementation, and Heutrix AI Guardrails for a scoped governance engagement. Direct entry is permitted when prerequisites are satisfied. No compulsory Diagnostics purchase, automatic fee credit or public pricing is introduced.',
      18:'The former pack contains retired service names, the unapproved AUD 950 Diagnostics fee, fixed scope/timing assumptions and unverified deletion commitments. The v2 review masters replace those defaults with explicit approval fields and retain the current Transformation handover and 30-day in-scope defect position.',
      19:'\tThe website approval is separate. The owner approved its privacy policy and terms on 12 September, waived independent review for that website launch and published them on 13 September. Shashane is the enquiry handler, Rochelle the backup. The waiver does not release client contracts. Document 07 identifies the approved public versions.',
      20:'\tThe former v1 native files and ZIP remain unchanged historical drafts. This package is a separate review revision. Do not issue either review ZIP to a client as a contract or edit a signed client copy informally.',
      21:'Analysis',22:'The document-use matrix in README identifies the forms and release gates. The operating procedures are in document 08; the single approval worksheet is document 09.',
      23:'Every paid project uses the approved Client Services Agreement plus a completed, signed Diagnostics Order Form or Statement of Work. Data access uses form A; AI uses form B; changes use form C; acceptance uses form D; handover and access closure use form E. Forms F and G support privacy requests and incidents.',
      24:'\tThe proposed revisions require explicit acceptance and a selected IP model. Liability caps, payment periods, confidentiality periods and other commercial choices remain marked for decision. These are drafting proposals until Australian legal review and owner approval are recorded for the exact versions.',
      25:'\tThe SOPs cover launch checks, enquiries, unsolicited sensitive information, issue/signing, data and suppliers, AI, changes and defects, closeout, incidents, privacy requests and marketing permissions. They keep live client records outside Git and distinguish the existing enquiry pathway from unapproved client systems.',
      26:'The Australian source review identifies consumer-guarantee and unfair-contract-term issues, federal privacy applicability, Victorian health-information handling, breach assessment and spam consent. These are review topics, not certifications of compliance. Document 09 supplies the primary source links and the questions for the legal adviser.',
      27:'\tMonday enquiries can follow the already approved website route and current information boundaries. Paid onboarding waits for the relevant approvals, signed scope and terms, payment/start conditions, insurance and safe data/access arrangements. A deadline, successful export or owner instruction to prepare a pack does not prove those conditions have been met.',
      28:'Conclusion',29:'The address and authorised directors are now populated from the owner decision. Complete document 09 with the formal notice email and execution method; insurance and commercial decisions explicitly remain pending. Resolve delivery rules, client-system/data approvals and legal review. Record affected-user and owner approval, effective date and exact versions in the legal review register. Then regenerate and check a separate approved issue bundle. Until that decision, use this pack for internal preparation and review only.'
    }
    for idx,txt in text_by_index.items():
        p=ps[idx]; texts=p.findall('.//w:t',ns)
        if texts:
            texts[0].text=txt
            for t in texts[1:]:t.text=''
        else:
            r=etree.SubElement(p,qn('w:r'));etree.SubElement(r,qn('w:t')).text=txt
    parts['word/document.xml']=etree.tostring(tree,xml_declaration=True,encoding='UTF-8',standalone=True)
    for n in list(parts):
        if re.match(r'word/header\d+\.xml$',n):
            t=etree.fromstring(parts[n]);ts=t.findall('.//w:t',ns)
            if ts:
                ts[0].text='INTERNAL REVIEW DRAFT'
                for x in ts[1:]:x.text=''
            parts[n]=etree.tostring(t,xml_declaration=True,encoding='UTF-8',standalone=True)
    with ZipFile(OUT/'00_Heutrix_Launch_Readiness_Memo.docx','w',ZIP_DEFLATED) as z:
        for n,b in parts.items():z.writestr(n,b)
    with ZipFile(REF) as z:
        unchanged=[n for n in z.namelist() if n!='word/document.xml' and not re.match(r'word/header\d+\.xml$',n)]
        assert all(z.read(n)==parts[n] for n in unchanged)

memo()
names={1:'01_Heutrix_Client_Services_Agreement.docx',2:'02_Heutrix_Diagnostics_Order_Form.docx',3:'03_Heutrix_Statement_of_Work.docx',4:'04_Heutrix_Mutual_Confidentiality_Agreement.docx',5:'05_Heutrix_Project_Control_Forms.docx'}
for n,name in names.items():reconcile_native(n,name)
for n,stem in [(6,'Project_Privacy_and_Notices'),(7,'Website_Policy_Reference'),(8,'Launch_SOPs'),(9,'Approval_Decisions')]:
    new_document(PACK/'source'/f'{n:02}_{stem}.md',f'{n:02}_Heutrix_{stem}.docx',f'{n:02}',n in (7,8,9))
(PACK/'source'/'reconciliation-changes.json').write_text(json.dumps(changes,ensure_ascii=False,indent=2),encoding='utf-8')
print(f'Built {len(list(OUT.glob("*.docx")))} review documents; {len(changes)} source replacements recorded.')
